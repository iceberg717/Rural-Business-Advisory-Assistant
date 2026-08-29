"""
Step 1: Clean & normalize all Mehsana datasets into consistent tables.
Joins Agriculture <-> Population via normalized village name (with dupe flagging).
Parses MSME's nested Activities JSON into flat columns.
Normalizes taluka/district spelling across financial_data.csv.
"""
import csv, json, re
import openpyxl

def norm(name):
    if name is None:
        return ""
    n = str(name).strip().lower()
    n = re.sub(r"\(.*?\)", "", n)          # drop parenthetical suffixes e.g. "Sardarpur (Cheekna)"
    n = re.sub(r"[^a-z0-9]+", " ", n).strip()
    return n

# ---------- 1. POPULATION (village level) ----------
wb = openpyxl.load_workbook("1787887989602_Mahesana_Population.xlsx", read_only=True)
ws = wb["EB-2404"]
rows = list(ws.iter_rows(values_only=True))
header = rows[0]
idx = {h: i for i, h in enumerate(header)}

pop_villages = []          # one row per village (TRU='Rural', villages are all rural here)
current_sub = None
for r in rows[1:]:
    level, name, tru = r[idx["Level"]], r[idx["Name"]], r[idx["TRU"]]
    if level == "SUB-DISTRICT" and tru == "Total":
        current_sub = name
    if level == "VILLAGE":
        pop_villages.append({
            "village_name": name,
            "village_name_norm": norm(name),
            "taluka": current_sub,
            "total_p": r[idx["TOT_P"]], "total_m": r[idx["TOT_M"]], "total_f": r[idx["TOT_F"]],
            "no_hh": r[idx["No_HH"]],
            "p_lit": r[idx["P_LIT"]], "m_lit": r[idx["M_LIT"]], "f_lit": r[idx["F_LIT"]],
            "p_sc": r[idx["P_SC"]], "p_st": r[idx["P_ST"]],
            "tot_work_p": r[idx["TOT_WORK_P"]], "tot_work_m": r[idx["TOT_WORK_M"]], "tot_work_f": r[idx["TOT_WORK_F"]],
            "main_cl_p": r[idx["MAIN_CL_P"]],   # cultivators
            "main_al_p": r[idx["MAIN_AL_P"]],   # agri labourers
        })

# flag name collisions across talukas
from collections import Counter
name_counts = Counter(v["village_name_norm"] for v in pop_villages)
dupe_names = {n for n, c in name_counts.items() if c > 1}

with open("population_clean.csv", "w", newline="") as f:
    fieldnames = list(pop_villages[0].keys()) + ["name_is_ambiguous"]
    w = csv.DictWriter(f, fieldnames=fieldnames)
    w.writeheader()
    for v in pop_villages:
        v["name_is_ambiguous"] = v["village_name_norm"] in dupe_names
        w.writerow(v)

pop_lookup = {}
for v in pop_villages:
    key = v["village_name_norm"]
    pop_lookup.setdefault(key, []).append(v)

print(f"Population: {len(pop_villages)} villages across {len(set(v['taluka'] for v in pop_villages))} talukas, "
      f"{len(dupe_names)} name collisions")

# ---------- 2. AGRICULTURE (village level) ----------
with open("1787887989601_Agricultural_Data.csv") as f:
    agri_rows = list(csv.DictReader(f))

matched, unmatched, ambiguous = 0, 0, 0
agri_out = []
for r in agri_rows:
    vname = r["VILL_NAME"]
    key = norm(vname)
    hits = pop_lookup.get(key, [])
    if len(hits) == 1:
        taluka = hits[0]["taluka"]
        matched += 1
    elif len(hits) > 1:
        taluka = "AMBIGUOUS:" + "/".join(sorted(set(h["taluka"] for h in hits)))
        ambiguous += 1
    else:
        taluka = "UNMATCHED"
        unmatched += 1
    agri_out.append({
        "village_name": vname,
        "village_name_norm": key,
        "taluka": taluka,
        "near_town": r["NEAR_TOWN"],
        "area_ha": r["AREA"],
        "tot_hh": r["T_HH"], "tot_p": r["T_P"], "tot_m": r["T_M"], "tot_f": r["T_F"],
        "main_crop_1": r["MAN_COMM1"], "main_crop_2": r["MAN_COMM2"], "main_crop_3": r["MAN_COMM3"],
        "tot_income": r["TOT_INC"], "tot_expense": r["TOT_EXP"],
        "tot_irrigated": r["TOT_IRR"], "un_irrigated": r["UN_IRR"],
        "bank_facility": r["BANK_FAC"], "comm_bank": r["COMM_BANK"], "coop_bank": r["COOP_BANK"],
        "power_supply": r["POWER_SUPL"], "power_agri": r["POWER_AGR"],
    })

with open("agriculture_clean.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(agri_out[0].keys()))
    w.writeheader()
    w.writerows(agri_out)

print(f"Agriculture: {len(agri_out)} villages -> matched {matched}, ambiguous {ambiguous}, unmatched {unmatched}")

# ---------- 3. MSME (unnest Activities JSON) ----------
wb2 = openpyxl.load_workbook("1787887989603_Mehsana_MSME_Data__1_.xlsx", read_only=True)
ws2 = wb2["Sheet1"]
rows2 = ws2.iter_rows(values_only=True)
header2 = next(rows2)
idx2 = {h: i for i, h in enumerate(header2)}

msme_out = []
bad_json = 0
ncols = len(header2)
for r in rows2:
    if len(r) < ncols:
        r = list(r) + [None] * (ncols - len(r))
    act_raw = r[idx2["Activities"]]
    nic_codes, descs = [], []
    try:
        acts = json.loads(act_raw) if act_raw else []
        for a in acts:
            nic_codes.append(str(a.get("NIC5DigitId", "")))
            descs.append(a.get("Description", ""))
    except Exception:
        bad_json += 1
    msme_out.append({
        "enterprise_name": r[idx2["EnterpriseName"]],
        "district": r[idx2["District"]],
        "pincode": r[idx2["Pincode"]],
        "registration_date": r[idx2["RegistrationDate"]],
        "address": r[idx2["CommunicationAddress"]],
        "nic_codes": ";".join(nic_codes),
        "activity_descriptions": ";".join(descs),
    })

with open("msme_clean.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(msme_out[0].keys()))
    w.writeheader()
    w.writerows(msme_out)

print(f"MSME: {len(msme_out)} enterprises unnested, {bad_json} rows with bad/missing Activities JSON")

# ---------- 4. FINANCIAL / SCHEME DATA (normalize taluka spelling) ----------
with open("1787887989602_financial_data.csv") as f:
    fin_rows = list(csv.DictReader(f))

for r in fin_rows:
    r["taluka_norm"] = norm(r["taluka"])
    r["district_norm"] = norm(r["district"])

with open("financial_clean.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(fin_rows[0].keys()))
    w.writeheader()
    w.writerows(fin_rows)

print(f"Financial/scheme data: {len(fin_rows)} rows normalized")

# ---------- 5. MARKET PRICE + LITERACY JSON -> flat CSVs ----------
with open("1787887989604_market_wise_price_report.json") as f:
    price_data = json.load(f)

with open("market_prices_clean.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["commodity_group", "commodity", "msp_rs_per_quintal", "price_rs_per_quintal", "price_date"])
    w.writeheader()
    for c in price_data["commodities"]:
        date_key = list(c["price_rs_per_quintal"].keys())[0]
        w.writerow({
            "commodity_group": c["commodity_group"],
            "commodity": c["commodity"],
            "msp_rs_per_quintal": c["msp_rs_per_quintal_2026_27"],
            "price_rs_per_quintal": c["price_rs_per_quintal"][date_key],
            "price_date": price_data["report_date"],
        })
print("Market prices: 15 commodities flattened")

with open("1787887989604_mehsana_district_literacy_2011.json") as f:
    lit = json.load(f)

with open("literacy_summary_clean.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["area_type", "population", "male_population", "female_population", "literate", "literacy_rate_percent"])
    w.writerow(["district_total", lit["district_total"]["population"], lit["district_total"]["male_population"],
                lit["district_total"]["female_population"], lit["district_total"]["literate"], lit["district_total"]["literacy_rate_percent"]])
    w.writerow(["rural", lit["rural"]["population"], lit["rural"]["male_population"], lit["rural"]["female_population"],
                lit["rural"]["literate"]["total"], lit["rural"]["literacy_rate_percent"]["total"]])
    w.writerow(["urban", lit["urban"]["population"], lit["urban"]["male_population"], lit["urban"]["female_population"],
                lit["urban"]["literate"], ""])
print("Literacy summary flattened")

print("\nAll cleaned files written to /home/claude/work/")
